import csv
import re
from decimal import Decimal
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from sensors.models import Sensor


COORDINATE_IN_DEVICE_NAME = re.compile(
    r'-N(?P<lat>\d+(?:\.\d+)?)-E(?P<lon>\d+(?:\.\d+)?)'
)
NUMBER = re.compile(r'\d{1,2}[,.]\d+')
WHAT3WORDS = re.compile(r'///[^\s"]+')


HEADER_MAP = {
    'standortbezeichnung': 'device_name',
    'device_name': 'device_name',
    'devicename': 'device_name',
    'kommune': 'operator_name',
    'operator_name': 'operator_name',
    'name': 'display_name',
    'display_name': 'display_name',
    'standort (adresse)': 'location_description',
    'standort_adresse': 'location_description',
    'location_description': 'location_description',
    'standort (gps/ w3w)': 'coordinates',
    'standort (gps/w3w)': 'coordinates',
    'coordinates': 'coordinates',
}

HEADER_ROW_SCAN_LIMIT = 10
KNOWN_HEADER_LABELS = frozenset(HEADER_MAP.keys())


class Command(BaseCommand):
    help = 'Delete all existing sensor metadata and import fresh data from XLSX/CSV/TSV.'

    def add_arguments(self, parser):
        parser.add_argument(
            'path',
            help='Path to XLSX/CSV/TSV file.',
        )
        parser.add_argument(
            '--delimiter',
            default=None,
            help='Column delimiter for CSV/TSV files. '
                 'Default: tab for .tsv, comma for .csv.',
        )

    def handle(self, *args, **options):
        path = Path(options['path'])

        if not path.exists():
            raise CommandError(f'File not found: {path}')

        created = 0
        skipped = 0

        # Alle Daten zunächst einlesen und validieren.
        rows = list(self._rows(path, options['delimiter']))

        # Erst wenn die Datei gelesen werden konnte, wird die Tabelle
        # geleert und anschließend komplett neu aufgebaut.
        with transaction.atomic():
            deleted, _ = Sensor.objects.all().delete()

            self.stdout.write(
                f'Deleted {deleted} existing sensor records.'
            )

            for row_number, row in rows:
                if not row.get('device_name'):
                    skipped += 1
                    self.stderr.write(
                        f'Row {row_number}: '
                        f'missing Standortbezeichnung/device_name'
                    )
                    continue

                device_name = row.get('device_name', '').strip()
                if self._is_header_label(device_name):
                    skipped += 1
                    self.stderr.write(
                        f'Row {row_number}: skipped header row '
                        f'({device_name!r})'
                    )
                    continue

                operator_name = row.get('operator_name', '').strip()
                display_name = row.get('display_name', '').strip()
                location_description = row.get(
                    'location_description', ''
                ).strip()
                coordinate_text = row.get('coordinates', '').strip()

                latitude, longitude = self._coordinates(
                    device_name,
                    coordinate_text,
                )

                defaults = {
                    'name': display_name or device_name,
                    'device_name': device_name,
                    'operator_name': operator_name,
                    'display_name': display_name,
                    'location_description': location_description,
                    'what3words': self._what3words(coordinate_text),
                    'latitude': latitude,
                    'longitude': longitude,
                    'sensor_type': Sensor.TYPE_STANDARD,
                    'active': True,
                }

                Sensor.objects.create(**defaults)
                created += 1

        self.stdout.write(
            self.style.SUCCESS(
                f'Sensor metadata import done: '
                f'{created} created, '
                f'{skipped} skipped, '
                f'{deleted} old records deleted.'
            )
        )

    def _rows(self, path: Path, delimiter: str | None):
        if path.suffix.lower() == '.xlsx':
            yield from self._xlsx_rows(path)
            return

        yield from self._csv_rows(path, delimiter)

    def _csv_rows(self, path: Path, delimiter: str | None):
        resolved_delimiter = delimiter

        if resolved_delimiter is None:
            resolved_delimiter = (
                ',' if path.suffix.lower() == '.csv' else '\t'
            )

        with path.open(
            encoding='utf-8-sig',
            newline='',
        ) as file:
            reader = csv.reader(
                file,
                delimiter=resolved_delimiter,
            )
            rows = list(reader)

        yield from self._dict_rows(rows)

    def _xlsx_rows(self, path: Path):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError(
                'openpyxl is required for .xlsx imports. '
                'Run: pip install openpyxl'
            ) from exc

        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

        sheet = workbook.active

        rows = [
            [
                cell if cell is not None else ''
                for cell in row
            ]
            for row in sheet.iter_rows(values_only=True)
        ]

        yield from self._dict_rows(rows)

    def _dict_rows(self, rows):
        if not rows:
            return

        header_index, mapped_headers = self._locate_header_row(rows)
        has_known_header = header_index is not None

        if has_known_header:
            data_rows = rows[header_index + 1:]
            data_start_row = header_index + 2
        else:
            mapped_headers = [
                'device_name',
                'operator_name',
                'display_name',
                'location_description',
                'coordinates',
            ]
            data_rows = rows
            data_start_row = 1

        for row_number, row in enumerate(
            data_rows,
            start=data_start_row,
        ):
            if not row:
                continue

            if not ''.join(
                str(value or '')
                for value in row
            ).strip():
                continue

            values = [
                str(value or '').strip()
                for value in row
            ]

            yield row_number, {
                header: values[index]
                if index < len(values)
                else ''
                for index, header in enumerate(mapped_headers)
                if header
            }

    def _normalize_header(self, value) -> str:
        return re.sub(
            r'\s+',
            ' ',
            str(value or '').strip().lower(),
        )

    def _is_header_label(self, value: str) -> bool:
        return self._normalize_header(value) in KNOWN_HEADER_LABELS

    def _locate_header_row(self, rows):
        for index, row in enumerate(rows[:HEADER_ROW_SCAN_LIMIT]):
            normalized = [
                self._normalize_header(value)
                for value in row
            ]
            mapped = [
                HEADER_MAP.get(header, '')
                for header in normalized
            ]
            if 'device_name' in mapped:
                return index, mapped
        return None, None

    def _coordinates(
        self,
        device_name: str,
        coordinate_text: str,
    ):
        match = COORDINATE_IN_DEVICE_NAME.search(device_name)

        if match:
            return (
                Decimal(match.group('lat')),
                Decimal(match.group('lon')),
            )

        numbers = [
            value.replace(',', '.')
            for value in NUMBER.findall(coordinate_text)
        ]

        if len(numbers) >= 2:
            return (
                Decimal(numbers[0]),
                Decimal(numbers[1]),
            )

        return None, None

    def _what3words(self, coordinate_text: str) -> str:
        match = WHAT3WORDS.search(coordinate_text)

        return match.group(0) if match else ''
